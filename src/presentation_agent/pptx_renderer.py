"""The only module permitted to import or manipulate python-pptx objects."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.util import Inches, Pt


def render(composition, path):
    if composition.get('units') != {'geometry': 'inches', 'font_size': 'points'}:
        raise ValueError('Unsupported composition units')
    prs = Presentation()
    prs.core_properties.author = 'Presentation Agent'
    prs.core_properties.last_modified_by = 'Presentation Agent'
    prs.core_properties.title = 'Synthetic pilot evaluation'
    prs.core_properties.subject = 'Milestone 1 offline fixture'
    prs.core_properties.comments = ''
    prs.slide_width = Inches(composition['width'])
    prs.slide_height = Inches(composition['height'])
    theme = composition['theme']
    for spec in composition['slides']:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme['colors']['background'])
        for element in spec['elements']:
            box = [Inches(v) for v in element['box']]
            if element['kind'] == 'text':
                shape = slide.shapes.add_textbox(*box)
                shape.name = element['id']
                frame = shape.text_frame
                frame.word_wrap = True
                frame.margin_left = frame.margin_right = 0
                frame.margin_top = frame.margin_bottom = 0
                for i, line in enumerate(element['text'].split('\n')):
                    paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    paragraph.text = line
                    paragraph.font.name = element['font']
                    paragraph.font.size = Pt(element['size'])
                    paragraph.font.color.rgb = RGBColor.from_string(element['color'])
                    paragraph.space_after = Pt(element['paragraph_space_after'])
                    paragraph.font.bold = element.get('bold', False) or (element['bold_first_paragraph'] and i == 0)
            elif element['kind'] == 'rule':
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box)
                shape.name = element['id']
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(element['color'])
                shape.line.fill.background()
            else:
                style = element['style']
                if style['type'] != 'column':
                    raise ValueError('Only column charts are supported in Milestone 1')
                data = CategoryChartData()
                data.categories = element['categories']
                data.add_series(element['series_name'], element['values'])
                shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *box, data)
                shape.name = element['id']
                chart = shape.chart
                chart.has_legend = False
                chart.has_title = True
                chart.chart_title.text_frame.text = element['title'] + ' (' + element['unit'] + ')'
                chart.font.name = element['font']
                chart.font.size = Pt(style['axis_size'])
                for p in chart.chart_title.text_frame.paragraphs:
                    p.font.name = element['font']
                    p.font.size = Pt(style['title_size'])
                axis = chart.value_axis
                axis.minimum_scale = style['axis_min']
                axis.maximum_scale = style['axis_max']
                axis.major_unit = style['axis_step']
                axis.major_tick_mark = XL_TICK_MARK.NONE
                axis.minor_tick_mark = XL_TICK_MARK.NONE
                axis.has_major_gridlines = True
                axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(style['grid_color'])
                axis.major_gridlines.format.line.width = Pt(.5)
                axis.format.line.fill.background()
                chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
                chart.category_axis.minor_tick_mark = XL_TICK_MARK.NONE
                chart.category_axis.format.line.color.rgb = RGBColor.from_string(style['axis_color'])
                chart.category_axis.format.line.width = Pt(.5)
                axis.tick_labels.font.name = element['font']
                chart.category_axis.tick_labels.font.name = element['font']
                chart.category_axis.tick_labels.font.size = Pt(style['axis_size'])
                plot = chart.plots[0]
                plot.gap_width = style['gap_width']
                plot.has_data_labels = True
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                plot.data_labels.font.name = element['font']
                plot.data_labels.font.size = Pt(style['label_size'])
                plot.data_labels.number_format = style['label_format']
                for point, color in zip(chart.series[0].points, style['point_colors']):
                    point.format.line.fill.background()
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = RGBColor.from_string(color)
    prs.save(str(path))


def inspect_pptx(path):
    """Return plain JSON-compatible observations from a reopened file."""
    prs = Presentation(str(path))
    slides = []
    for slide in prs.slides:
        elements = []
        for shape in slide.shapes:
            item = {'id': shape.name, 'box': [shape.left / 914400, shape.top / 914400, shape.width / 914400, shape.height / 914400]}
            if shape.has_chart:
                chart = shape.chart
                item.update(kind='chart', categories=[c.label for c in chart.plots[0].categories], values=list(chart.series[0].values), embedded_workbook=chart.part.chart_workbook.xlsx_part is not None, font=chart.font.name)
                item.update(series_name=chart.series[0].name,
                            title=chart.chart_title.text_frame.text,
                            chart_type=int(chart.chart_type), axis_min=chart.value_axis.minimum_scale,
                            axis_max=chart.value_axis.maximum_scale, axis_step=chart.value_axis.major_unit,
                            label_format=chart.plots[0].data_labels.number_format,
                            workbook_cells=workbook_cells(chart.part.chart_workbook.xlsx_part.blob))
            elif shape.shape_type == 1:
                item.update(kind='rule', color=str(shape.fill.fore_color.rgb))
            elif shape.has_text_frame:
                item.update(kind='text', text=shape.text, fonts=[p.font.name for p in shape.text_frame.paragraphs], sizes=[p.font.size.pt for p in shape.text_frame.paragraphs], bold=[p.font.bold for p in shape.text_frame.paragraphs])
            else:
                item['kind'] = 'other'
            elements.append(item)
        slides.append(elements)
    return {'width': prs.slide_width / 914400, 'height': prs.slide_height / 914400, 'slides': slides}


def workbook_cells(blob):
    """Inspect actual OOXML workbook cells, independently of the chart cache."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(BytesIO(blob)) as z:
        strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            strings = [''.join(n.itertext()) for n in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('s:si', ns)]
        sheet = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        cells = {}
        for cell in sheet.findall('.//s:c', ns):
            value = cell.find('s:v', ns)
            if value is not None:
                cells[cell.attrib['r']] = strings[int(value.text)] if cell.get('t') == 's' else float(value.text)
        return cells
